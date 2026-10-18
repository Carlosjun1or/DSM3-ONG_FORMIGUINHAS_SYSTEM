import csv
import io
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from django import forms
from django.contrib import messages
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone

from acao.models import Acao, AcaoParticipante
from doacao.models import Doacao, DoacaoFinanceira, ItemDoacaoMaterial
from estoque.models import Material, MovimentacaoEstoque
from praia.models import Praia
from residuo.models import ColetaResiduo, ItemColetaResiduo, ResiduoColetado
from usuario.models import Usuario, Voluntario
from usuario.views import render_acesso_bloqueado


RELATORIOS = [
    ('geral', 'Relatório geral'),
    ('indicadores', 'Indicadores por período / praia'),
    ('acoes', 'Relatório de ações'),
    ('voluntarios', 'Relatório de voluntários'),
    ('presenca', 'Relatório de presença'),
    ('residuos', 'Relatório de resíduos'),
    ('estoque', 'Relatório de estoque'),
    ('doacoes', 'Relatório de doações'),
]
PERIODOS = [('MES', 'Mês atual'), ('TRIMESTRE', 'Trimestre atual'), ('ANO', 'Ano atual')]


class RelatorioFiltroForm(forms.Form):
    relatorio = forms.ChoiceField(label='Relatório', choices=RELATORIOS)
    data_inicio = forms.DateField(
        label='Data inicial',
        required=False,
        widget=forms.DateInput(attrs={'type': 'date'}),
    )
    data_fim = forms.DateField(
        label='Data final',
        required=False,
        widget=forms.DateInput(attrs={'type': 'date'}),
    )
    periodo = forms.ChoiceField(label='Período do painel', choices=PERIODOS, required=False)
    praia = forms.ModelChoiceField(
        label='Praia / local',
        queryset=Praia.objects.none(),
        required=False,
        empty_label='Todas as praias',
    )
    acao = forms.ModelChoiceField(
        label='Ação',
        queryset=Acao.objects.none(),
        required=False,
        empty_label='Todas as ações',
    )
    tipo_acao = forms.ChoiceField(
        label='Tipo de ação',
        choices=[('', 'Todos os tipos')] + list(Acao.TIPO_CHOICES),
        required=False,
    )
    tipo_residuo = forms.ChoiceField(
        label='Tipo de resíduo',
        choices=[('', 'Todos os tipos')] + list(ResiduoColetado.TIPO_CHOICES),
        required=False,
    )

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        hoje = timezone.localdate()
        if not self.is_bound:
            self.initial.setdefault('relatorio', 'geral')
            self.initial.setdefault('data_inicio', hoje.replace(day=1))
            self.initial.setdefault('data_fim', hoje)
            self.initial.setdefault('periodo', 'MES')
        self.fields['praia'].queryset = Praia.objects.order_by('nome')
        self.fields['acao'].queryset = Acao.objects.select_related('praia').order_by(
            '-data', '-id_acao'
        )
        for field in self.fields.values():
            field.widget.attrs.setdefault('class', 'form-control')

    def clean(self):
        dados = super().clean()
        inicio = dados.get('data_inicio')
        fim = dados.get('data_fim')
        if inicio and fim and inicio > fim:
            self.add_error('data_fim', 'A data final deve ser igual ou posterior à data inicial.')
        return dados


def _periodo_efetivo(dados):
    hoje = timezone.localdate()
    if dados['relatorio'] != 'indicadores':
        return dados.get('data_inicio') or date(1900, 1, 1), dados.get('data_fim') or hoje

    periodo = dados.get('periodo') or 'MES'
    if periodo == 'ANO':
        return hoje.replace(month=1, day=1), hoje
    if periodo == 'TRIMESTRE':
        mes_inicial = ((hoje.month - 1) // 3) * 3 + 1
        return hoje.replace(month=mes_inicial, day=1), hoje
    return hoje.replace(day=1), hoje


def _presenca_filter():
    return Q(compareceu=True) | Q(status=AcaoParticipante.STATUS_PRESENTE)


def _residuos_no_periodo(inicio, fim, tipo='', praia_id=None, acao_id=None):
    linhas = []
    itens_registrados = ItemColetaResiduo.objects.filter(
        coleta__data_coleta__date__range=(inicio, fim),
    ).select_related(
        'coleta__acao__praia',
        'coleta__bag__condominio__praia',
    ).order_by('coleta__data_coleta', 'id_item')
    if tipo:
        itens_registrados = itens_registrados.filter(tipo=tipo)
    if praia_id:
        itens_registrados = itens_registrados.filter(
            Q(coleta__acao__praia_id=praia_id)
            | Q(coleta__bag__condominio__praia_id=praia_id)
        )
    if acao_id:
        itens_registrados = itens_registrados.filter(coleta__acao_id=acao_id)

    for item in itens_registrados:
        coleta = item.coleta
        acao = coleta.acao
        bag = coleta.bag
        linhas.append([
            timezone.localtime(coleta.data_coleta).strftime('%d/%m/%Y'),
            'Ação' if coleta.origem == 'ACAO' else 'Bag',
            str(acao) if acao else '-',
            bag.codigo if bag else '-',
            (acao.praia.nome if acao else bag.condominio.praia.nome if bag and bag.condominio.praia else '-'),
            item.get_tipo_display() or 'Não informado',
            item.get_categoria_display() or 'Não informada',
            str(item.quantidade or 0),
            item.get_unidade_display() or 'Não informada',
        ])

    residuos_acao = ResiduoColetado.objects.filter(
        dt_registro__date__range=(inicio, fim),
    ).select_related('acao__praia').order_by('dt_registro', 'id_residuo')
    if tipo:
        residuos_acao = residuos_acao.filter(tipo=tipo)
    if praia_id:
        residuos_acao = residuos_acao.filter(acao__praia_id=praia_id)
    if acao_id:
        residuos_acao = residuos_acao.filter(acao_id=acao_id)
    for residuo in residuos_acao:
        linhas.append([
            timezone.localtime(residuo.dt_registro).strftime('%d/%m/%Y'),
            'Ação',
            str(residuo.acao),
            '-',
            residuo.acao.praia.nome,
            residuo.get_tipo_display() or 'Não informado',
            residuo.get_categoria_display() or 'Não informada',
            str(residuo.quantidade or 0),
            residuo.get_unidade_display() or 'Não informada',
        ])
    return sorted(linhas, key=lambda linha: datetime.strptime(linha[0], '%d/%m/%Y'))


def _resumo_quantidades_por_unidade(linhas):
    totais = defaultdict(Decimal)
    for linha in linhas:
        quantidade = Decimal(linha[7] or 0)
        totais[linha[8]] += quantidade
    return ', '.join(f'{quantidade:g} {unidade}' for unidade, quantidade in sorted(totais.items())) or '0'


def _resumo_saldos(materiais):
    totais = defaultdict(int)
    for material in materiais:
        totais[material.get_unidade_display()] += material.saldo
    return ', '.join(f'{quantidade} {unidade}' for unidade, quantidade in sorted(totais.items())) or '0'



def _criar_relatorio(dados, inicio, fim):
    tipo = dados['relatorio']
    praia_id = dados.get('praia').pk if dados.get('praia') else None
    acao_id = dados.get('acao').pk if dados.get('acao') else None
    tipo_acao = dados.get('tipo_acao') or ''
    tipo_residuo = dados.get('tipo_residuo') or ''
    periodo_label = f'{inicio:%d/%m/%Y} a {fim:%d/%m/%Y}'

    relatorio = {
        'title': dict(RELATORIOS)[tipo],
        'period': periodo_label,
        'metrics': [],
        'sections': [],
        'attendance_action': dados.get('acao') if tipo == 'presenca' else None,
    }

    if tipo == 'geral':
        acoes = Acao.objects.filter(data__range=(inicio, fim))
        presencas = AcaoParticipante.objects.filter(
            acao__data__range=(inicio, fim),
            acao__status=Acao.STATUS_CONCLUIDA,
        ).filter(_presenca_filter())
        doacoes = Doacao.objects.filter(
            data_recebimento__range=(inicio, fim),
            status=Doacao.CONFIRMADA,
        )
        linhas_residuos = _residuos_no_periodo(inicio, fim, praia_id=praia_id)
        materiais = list(Material.objects.com_saldo().order_by('nome'))
        if praia_id:
            acoes = acoes.filter(praia_id=praia_id)
            presencas = presencas.filter(acao__praia_id=praia_id)
            doacoes = doacoes.filter(Q(praia_id=praia_id) | Q(acao__praia_id=praia_id))
            linhas_residuos = [linha for linha in linhas_residuos if linha[4] == dados['praia'].nome]

        status_acoes = list(
            acoes.values('status').annotate(total=Count('id_acao')).order_by('status')
        )
        linhas_status_acoes = [
            [dict(Acao.STATUS_CHOICES).get(item['status'], item['status']), item['total']]
            for item in status_acoes
        ]
        linhas_residuos_mes = []
        por_mes_unidade = defaultdict(Decimal)
        for linha in linhas_residuos:
            mes = datetime.strptime(linha[0], '%d/%m/%Y').strftime('%Y-%m')
            por_mes_unidade[(mes, linha[8])] += Decimal(linha[7] or 0)
        linhas_residuos_mes = [
            [f'{mes[5:7]}/{mes[:4]}', unidade, f'{quantidade:g}']
            for (mes, unidade), quantidade in sorted(por_mes_unidade.items())
        ]
        financeiro_por_doacao = {
            item['doacao_id']: item['valor']
            for item in DoacaoFinanceira.objects.filter(
                doacao__in=doacoes,
            ).values('doacao_id', 'valor')
        }
        total_financeiro = sum(financeiro_por_doacao.values(), Decimal('0'))
        itens_doacao = list(
            ItemDoacaoMaterial.objects.filter(doacao__in=doacoes).select_related('doacao')
        )
        linhas_doacoes = [
            ['Financeiras', doacoes.filter(tipo=Doacao.FINANCEIRA).count(), f'R$ {total_financeiro:.2f}'],
            ['Materiais', doacoes.filter(tipo=Doacao.MATERIAL).count(), len(itens_doacao)],
        ]
        relatorio['metrics'] = [
            ('Voluntários cadastrados', Voluntario.objects.count()),
            ('Voluntários ativos', Voluntario.objects.filter(status='ATIVO').count()),
            ('Ações no período', acoes.count()),
            ('Ações concluídas', acoes.filter(status=Acao.STATUS_CONCLUIDA).count()),
            ('Voluntários presentes', presencas.values('voluntario_id').distinct().count()),
            ('Resíduos coletados', _resumo_quantidades_por_unidade(linhas_residuos)),
            ('Bags coletadas', ColetaResiduo.objects.filter(
                origem='BAG',
                data_coleta__date__range=(inicio, fim),
                **({'bag__condominio__praia_id': praia_id} if praia_id else {}),
            ).count()),
            ('Itens com estoque baixo', sum(material.baixo_estoque for material in materiais)),
            ('Saldo atual em estoque', _resumo_saldos(materiais)),
            ('Doações confirmadas', doacoes.count()),
            ('Total financeiro doado', f'R$ {total_financeiro:.2f}'),
        ]
        relatorio['sections'] = [
            {
                'title': 'Ações por situação',
                'headers': ['Situação', 'Quantidade'],
                'rows': linhas_status_acoes,
            },
            {
                'title': 'Resíduos coletados por mês e unidade',
                'headers': ['Mês/ano', 'Unidade', 'Quantidade'],
                'rows': linhas_residuos_mes,
            },
            {
                'title': 'Resumo de doações confirmadas',
                'headers': ['Tipo', 'Quantidade de doações', 'Total'],
                'rows': linhas_doacoes,
            },
            {
                'title': 'Estoque disponível',
                'headers': ['Material', 'Código', 'Saldo', 'Unidade', 'Estoque mínimo', 'Situação'],
                'rows': [[
                    material.nome,
                    material.codigo,
                    material.saldo,
                    material.get_unidade_display(),
                    material.estoque_minimo,
                    'Baixo' if material.baixo_estoque else 'Normal',
                ] for material in materiais],
            },
        ]

    elif tipo == 'acoes':
        acoes = Acao.objects.filter(
            data__range=(inicio, fim),
            status=Acao.STATUS_CONCLUIDA,
        ).select_related('praia').order_by('-data', 'id_acao')
        if acao_id:
            acoes = acoes.filter(pk=acao_id)
        if tipo_acao:
            acoes = acoes.filter(tipo=tipo_acao)
        if praia_id:
            acoes = acoes.filter(praia_id=praia_id)

        linhas = []
        total_participantes = 0
        total_material = 0
        for acao in acoes:
            participantes = list(
                acao.participantes.filter(_presenca_filter()).select_related('voluntario')
                .order_by('voluntario__nome')
            )
            total_participantes += len(participantes)
            nomes = '; '.join(
                f'{p.voluntario.nome} ({p.voluntario.telefone})' for p in participantes
            ) or '-'
            movimentos = list(
                acao.movimentacoes_estoque.filter(tipo='SAIDA').select_related('material')
            )
            materiais = '; '.join(
                f'{mov.material.nome}: {mov.quantidade} {mov.material.get_unidade_display()}'
                for mov in movimentos
            ) or '-'
            total_material += len(movimentos)
            residuos = _residuos_no_periodo(inicio, fim, acao_id=acao.pk)
            linhas.append([
                acao.data.strftime('%d/%m/%Y'),
                acao.get_tipo_display(),
                acao.praia.nome,
                acao.local or '-',
                acao.get_status_display(),
                nomes,
                materiais,
                _resumo_quantidades_por_unidade(residuos),
            ])
        relatorio['metrics'] = [
            ('Ações concluídas', len(linhas)),
            ('Participações', total_participantes),
            ('Saídas de materiais vinculadas', total_material),
        ]
        relatorio['sections'] = [{
            'title': 'Ações e resultados',
            'headers': ['Data', 'Tipo', 'Praia', 'Local', 'Status', 'Participantes e telefones', 'Materiais utilizados', 'Resíduos coletados'],
            'rows': linhas,
        }]

    elif tipo == 'voluntarios':
        filtro_presenca = Q(acoes_participadas__compareceu=True) | Q(
            acoes_participadas__status=AcaoParticipante.STATUS_PRESENTE
        )
        acoes_filtro = Q(
            acoes_participadas__acao__data__range=(inicio, fim),
            acoes_participadas__acao__status=Acao.STATUS_CONCLUIDA,
        ) & filtro_presenca
        if tipo_acao:
            acoes_filtro &= Q(acoes_participadas__acao__tipo=tipo_acao)
        voluntarios = Voluntario.objects.annotate(
            participacoes=Count(
                'acoes_participadas',
                filter=acoes_filtro,
                distinct=True,
            )
        ).order_by('-participacoes', 'nome')
        total_cadastrados = Voluntario.objects.count()
        relatorio['metrics'] = [
            ('Voluntários cadastrados', total_cadastrados),
            ('Voluntários participantes no período', voluntarios.filter(participacoes__gt=0).count()),
            ('Participações no período', sum(v.participacoes for v in voluntarios)),
        ]
        relatorio['sections'] = [{
            'title': 'Participação dos voluntários',
            'headers': ['Nome', 'Telefone', 'Status cadastral', 'Ações frequentadas no período'],
            'rows': [[v.nome, v.telefone, v.get_status_display(), v.participacoes] for v in voluntarios],
        }]

    elif tipo == 'presenca':
        participantes = AcaoParticipante.objects.filter(
            acao__data__range=(inicio, fim),
        ).filter(_presenca_filter()).select_related(
            'acao__praia', 'voluntario'
        ).order_by('acao__data', 'acao__id_acao', 'voluntario__nome')
        if acao_id:
            participantes = participantes.filter(acao_id=acao_id)
        if tipo_acao:
            participantes = participantes.filter(acao__tipo=tipo_acao)
        if praia_id:
            participantes = participantes.filter(acao__praia_id=praia_id)
        por_acao = defaultdict(list)
        for participante in participantes:
            por_acao[participante.acao].append([
                participante.voluntario.nome,
                participante.voluntario.telefone,
                '________________________',
            ])
        relatorio['metrics'] = [
            ('Ações no período', len(por_acao)),
            ('Presenças registradas', sum(len(linhas) for linhas in por_acao.values())),
        ]
        relatorio['sections'] = [
            {
                'title': f'{acao} — {acao.data:%d/%m/%Y}',
                'headers': ['Nome completo', 'Telefone', 'Assinatura'],
                'rows': linhas,
            }
            for acao, linhas in por_acao.items()
        ] or [{
            'title': 'Lista de presença',
            'headers': ['Nome completo', 'Telefone', 'Assinatura'],
            'rows': [],
        }]

    elif tipo == 'residuos':
        linhas = _residuos_no_periodo(inicio, fim, tipo_residuo, praia_id, acao_id)
        por_mes = defaultdict(list)
        for linha in linhas:
            partes = linha[0].split('/')
            chave = f'{partes[2]}-{partes[1]}'
            por_mes[chave].append(linha)
        resumo_mensal = [
            [f'{mes[5:7]}/{mes[:4]}', unidade, f'{sum(Decimal(l[7] or 0) for l in itens if l[8] == unidade):g}']
            for mes, itens in sorted(por_mes.items())
            for unidade in sorted({l[8] for l in itens})
        ]
        relatorio['metrics'] = [
            ('Registros de coleta', len(linhas)),
            ('Quantidade por unidade', _resumo_quantidades_por_unidade(linhas)),
        ]
        relatorio['sections'] = [
            {
                'title': 'Quantidade coletada por mês e unidade',
                'headers': ['Mês/ano', 'Unidade', 'Quantidade total'],
                'rows': resumo_mensal,
            },
            {
                'title': 'Ações e bags coletadas',
                'headers': ['Data', 'Origem', 'Ação', 'Bag', 'Praia', 'Tipo', 'Categoria', 'Quantidade', 'Unidade'],
                'rows': linhas,
            },
        ]

    elif tipo == 'estoque':
        materiais = list(Material.objects.com_saldo().order_by('nome'))
        linhas_saldo = [[
            material.nome,
            material.codigo,
            material.categoria,
            material.saldo,
            material.get_unidade_display(),
            material.estoque_minimo,
            'Baixo' if material.baixo_estoque else 'Normal',
        ] for material in materiais]
        movimentos = MovimentacaoEstoque.objects.filter(
            dt_cadastro__date__range=(inicio, fim),
        ).select_related('material', 'acao').order_by('-dt_cadastro')
        linhas_movimentos = [[
            timezone.localtime(m.dt_cadastro).strftime('%d/%m/%Y %H:%M'),
            m.get_tipo_display(),
            m.material.nome,
            m.quantidade,
            m.material.get_unidade_display(),
            str(m.acao) if m.acao_id else '-',
            m.registrado_por_nome or '-',
            m.justificativa or '-',
        ] for m in movimentos]
        relatorio['metrics'] = [
            ('Itens cadastrados', len(materiais)),
            ('Quantidade total em estoque', _resumo_saldos(materiais)),
            ('Itens com estoque baixo', sum(m.baixo_estoque for m in materiais)),
            ('Movimentações no período', len(linhas_movimentos)),
        ]
        relatorio['sections'] = [
            {
                'title': 'Saldo atual e nível mínimo',
                'headers': ['Material', 'Código', 'Categoria', 'Saldo disponível', 'Unidade', 'Estoque mínimo', 'Situação'],
                'rows': linhas_saldo,
            },
            {
                'title': 'Histórico de entradas e saídas',
                'headers': ['Data/hora', 'Movimentação', 'Material', 'Quantidade', 'Unidade', 'Ação', 'Responsável', 'Justificativa'],
                'rows': linhas_movimentos,
            },
        ]

    elif tipo == 'indicadores':
        acoes = Acao.objects.filter(
            data__range=(inicio, fim),
            status=Acao.STATUS_CONCLUIDA,
        )
        presencas = AcaoParticipante.objects.filter(
            acao__data__range=(inicio, fim),
            acao__status=Acao.STATUS_CONCLUIDA,
            voluntario__status='ATIVO',
        ).filter(_presenca_filter())
        if praia_id:
            acoes = acoes.filter(praia_id=praia_id)
            presencas = presencas.filter(acao__praia_id=praia_id)
        linhas_residuos = _residuos_no_periodo(inicio, fim, praia_id=praia_id)
        materiais = list(Material.objects.com_saldo())
        praias = Praia.objects.order_by('nome')
        if praia_id:
            praias = praias.filter(pk=praia_id)
        linhas_praia = []
        for praia in praias:
            residuos_praia = [linha for linha in linhas_residuos if linha[4] == praia.nome]
            presencas_praia = presencas.filter(acao__praia=praia).values(
                'voluntario_id'
            ).distinct().count()
            linhas_praia.append([
                praia.nome,
                acoes.filter(praia=praia).count(),
                presencas_praia,
                len(residuos_praia),
                _resumo_quantidades_por_unidade(residuos_praia),
            ])
        relatorio['metrics'] = [
            ('Voluntários ativos participantes', presencas.values('voluntario_id').distinct().count()),
            ('Ações concluídas', acoes.count()),
            ('Registros de resíduos', len(linhas_residuos)),
            ('Materiais cadastrados', len(materiais)),
            ('Saldo em estoque', _resumo_saldos(materiais)),
        ]
        relatorio['sections'] = [{
            'title': 'Indicadores por praia',
            'headers': ['Praia', 'Ações concluídas', 'Voluntários ativos', 'Registros de resíduo', 'Quantidade coletada por unidade'],
            'rows': linhas_praia,
        }]

    elif tipo == 'doacoes':
        doacoes = Doacao.objects.filter(
            data_recebimento__range=(inicio, fim),
            status=Doacao.CONFIRMADA,
        ).select_related('doador').order_by('-data_recebimento', '-id_doacao')
        itens = ItemDoacaoMaterial.objects.filter(
            doacao__in=doacoes,
        ).select_related('doacao__doador').order_by('nome')
        valores_financeiros = {
            item['doacao_id']: item['valor']
            for item in DoacaoFinanceira.objects.filter(
                doacao__in=doacoes,
            ).values('doacao_id', 'valor')
        }
        total_financeiro = sum(valores_financeiros.values(), Decimal('0'))
        itens_por_doacao = defaultdict(list)
        materiais_por_nome = defaultdict(Decimal)
        materiais_por_doador = defaultdict(lambda: defaultdict(Decimal))
        for item in itens:
            doador = item.doacao.doador.nome if item.doacao.doador_id else 'Anônimo'
            itens_por_doacao[item.doacao_id].append(
                f'{item.nome}: {item.quantidade:g} {item.unidade}'
            )
            chave_material = (item.nome, item.unidade)
            materiais_por_nome[chave_material] += item.quantidade
            materiais_por_doador[doador][item.unidade] += item.quantidade
        linhas_doacoes = []
        contribuicoes = defaultdict(lambda: {'count': 0, 'financeiro': Decimal('0'), 'materiais': defaultdict(Decimal)})
        for doacao in doacoes:
            doador = doacao.doador.nome if doacao.doador_id else 'Anônimo'
            valor = valores_financeiros.get(doacao.pk, Decimal('0'))
            linhas_doacoes.append([
                doacao.data_recebimento.strftime('%d/%m/%Y'),
                doacao.get_tipo_display(),
                doador,
                f'R$ {valor:.2f}' if valor else '-',
                '; '.join(itens_por_doacao[doacao.pk]) or '-',
                doacao.get_finalidade_display(),
            ])
            contribuicoes[doador]['count'] += 1
            contribuicoes[doador]['financeiro'] += valor
        for item in itens:
            doador = item.doacao.doador.nome if item.doacao.doador_id else 'Anônimo'
            contribuicoes[doador]['materiais'][item.unidade] += item.quantidade
        linhas_materiais = [
            [nome, unidade, f'{quantidade:g}']
            for (nome, unidade), quantidade in sorted(materiais_por_nome.items())
        ]
        linhas_doadores = [
            [
                nome,
                resumo['count'],
                f'R$ {resumo["financeiro"]:.2f}',
                ', '.join(f'{q:g} {unidade}' for unidade, q in sorted(resumo['materiais'].items())) or '-',
            ]
            for nome, resumo in sorted(
                contribuicoes.items(),
                key=lambda item: (-item[1]['count'], -item[1]['financeiro'], item[0]),
            )
        ]
        relatorio['metrics'] = [
            ('Doações confirmadas', doacoes.count()),
            ('Total financeiro', f'R$ {total_financeiro:.2f}'),
            ('Materiais recebidos', len(itens)),
            ('Doadores identificados', doacoes.exclude(doador__isnull=True).values('doador_id').distinct().count()),
        ]
        relatorio['sections'] = [
            {
                'title': 'Doações confirmadas',
                'headers': ['Data', 'Tipo', 'Doador', 'Valor financeiro', 'Materiais e quantidades', 'Finalidade'],
                'rows': linhas_doacoes,
            },
            {
                'title': 'Materiais recebidos',
                'headers': ['Material', 'Unidade', 'Quantidade total'],
                'rows': linhas_materiais,
            },
            {
                'title': 'Doadores que mais contribuíram',
                'headers': ['Doador', 'Doações', 'Total financeiro', 'Materiais doados'],
                'rows': linhas_doadores,
            },
        ]

    return relatorio


def _resposta_excel(relatorio, tipo, hoje):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    resumo = workbook.active
    resumo.title = 'Resumo'
    resumo.append([relatorio['title']])
    resumo.append([f'Período: {relatorio["period"]}'])
    resumo.append([])
    resumo.append(['Indicador', 'Resultado'])
    for celula in resumo[4]:
        celula.font = Font(bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor='155E4B')
    for rotulo, valor in relatorio['metrics']:
        resumo.append([rotulo, str(valor)])
        resumo.cell(resumo.max_row, 2).data_type = 's'
    resumo.freeze_panes = 'A5'
    resumo.column_dimensions['A'].width = 38
    resumo.column_dimensions['B'].width = 36

    for indice, secao in enumerate(relatorio['sections'], start=1):
        nome = ''.join(char for char in secao['title'] if char not in '[]:*?/\\')[:31]
        planilha = workbook.create_sheet(nome or f'Relatório {indice}')
        planilha.append([relatorio['title']])
        planilha.append([f'Período: {relatorio["period"]}'])
        planilha.append([])
        planilha.append(secao['headers'])
        for celula in planilha[4]:
            celula.font = Font(bold=True, color='FFFFFF')
            celula.fill = PatternFill('solid', fgColor='155E4B')
        for linha in secao['rows']:
            planilha.append([
                str(valor) if valor is not None else ''
                for valor in linha
            ])
            for celula in planilha[planilha.max_row]:
                celula.data_type = 's'
        planilha.freeze_panes = 'A5'
        planilha.auto_filter.ref = f'A4:{planilha.cell(planilha.max_row, len(secao["headers"])).coordinate}'
        for coluna in planilha.columns:
            largura = min(max(len(str(c.value or '')) for c in coluna) + 2, 48)
            planilha.column_dimensions[coluna[0].column_letter].width = max(largura, 12)
    conteudo = io.BytesIO()
    workbook.save(conteudo)
    resposta = HttpResponse(
        conteudo.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resposta['Content-Disposition'] = (
        f'attachment; filename="{tipo}_{hoje.isoformat()}.xlsx"'
    )
    return resposta


def _resposta_pdf(relatorio, tipo, hoje):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        LongTable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.45 * inch,
        rightMargin=0.45 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle(
        'TituloRelatorio',
        parent=estilos['Title'],
        alignment=TA_CENTER,
        textColor=colors.HexColor('#155E4B'),
    )
    pequeno = ParagraphStyle(
        'TabelaRelatorio',
        parent=estilos['BodyText'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
    )
    elementos = [
        Paragraph(escape(relatorio['title']), titulo),
        Paragraph(f'Período: {escape(relatorio["period"])}', estilos['Normal']),
        Spacer(1, 12),
    ]
    if relatorio['attendance_action']:
        elementos.extend([
            Paragraph(f'Ação: {escape(str(relatorio["attendance_action"]))}', estilos['Heading2']),
            Paragraph(
                f'Data da ação: {relatorio["attendance_action"].data:%d/%m/%Y}',
                estilos['Normal'],
            ),
            Spacer(1, 8),
        ])
    if relatorio['metrics']:
        elementos.append(Paragraph('Resumo', estilos['Heading2']))
        resumo = Table(
            [
                [Paragraph('Indicador', pequeno), Paragraph('Resultado', pequeno)],
                *[
                    [Paragraph(escape(str(rotulo)), pequeno), Paragraph(escape(str(valor)), pequeno)]
                    for rotulo, valor in relatorio['metrics']
                ],
            ],
            colWidths=[documento.width * 0.38, documento.width * 0.62],
        )
        resumo.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#155E4B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#CFD8D3')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F7F4')]),
        ]))
        elementos.extend([resumo, Spacer(1, 12)])
    for indice, secao in enumerate(relatorio['sections']):
        elementos.append(Paragraph(escape(secao['title']), estilos['Heading2']))
        dados = [[Paragraph(escape(str(item)), pequeno) for item in secao['headers']]]
        linhas = secao['rows'] or [['Nenhum registro encontrado.'] + [''] * (len(secao['headers']) - 1)]
        dados.extend([
            [Paragraph(escape(str(item)), pequeno) for item in linha]
            for linha in linhas
        ])
        largura = documento.width / max(len(secao['headers']), 1)
        tabela = LongTable(
            dados,
            colWidths=[largura] * len(secao['headers']),
            repeatRows=1,
        )
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#155E4B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#CFD8D3')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F7F4')]),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]))
        elementos.extend([tabela, Spacer(1, 12)])
        if indice < len(relatorio['sections']) - 1 and len(secao['rows']) > 35:
            elementos.append(PageBreak())
    documento.build(elementos)
    resposta = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    resposta['Content-Disposition'] = f'attachment; filename="{tipo}_{hoje.isoformat()}.pdf"'
    return resposta


def _resposta_csv(relatorio, tipo, hoje):
    resposta = HttpResponse(content_type='text/csv; charset=utf-8')
    resposta['Content-Disposition'] = f'attachment; filename="{tipo}_{hoje.isoformat()}.csv"'
    resposta.write('\ufeff')
    escritor = csv.writer(resposta, delimiter=';')
    for secao in relatorio['sections']:
        escritor.writerow([secao['title']])
        escritor.writerow(secao['headers'])
        for linha in secao['rows']:
            seguro = [
                "'" + str(valor) if str(valor).startswith(('=', '+', '-', '@')) else valor
                for valor in linha
            ]
            escritor.writerow(seguro)
        escritor.writerow([])
    return resposta


def relatorios_view(request):
    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, 'Acesso restrito: faça login para acessar os relatórios.')
        return redirect('login')
    try:
        usuario = Usuario.objects.select_related('id_voluntario').get(id_usuario=usuario_id)
    except Usuario.DoesNotExist:
        request.session.flush()
        return redirect('login')
    if usuario.tipo != 'ADMIN':
        return render_acesso_bloqueado(request)

    parametros = request.GET.copy()
    hoje = timezone.localdate()
    parametros.setdefault('relatorio', 'indicadores')
    parametros.setdefault('periodo', 'MES')
    parametros.setdefault('data_inicio', hoje.replace(day=1).isoformat())
    parametros.setdefault('data_fim', hoje.isoformat())
    form = RelatorioFiltroForm(parametros)
    if not form.is_valid():
        return render(request, 'sistema/relatorios.html', {
            'usuario': usuario,
            'form': form,
            'report': None,
            'report_types': RELATORIOS,
        }, status=400)

    dados = form.cleaned_data
    inicio, fim = _periodo_efetivo(dados)
    report = _criar_relatorio(dados, inicio, fim)
    formato = request.GET.get('formato', '')
    tipo = dados['relatorio']
    hoje = timezone.localdate()
    if formato == 'xlsx':
        return _resposta_excel(report, tipo, hoje)
    if formato == 'pdf':
        return _resposta_pdf(report, tipo, hoje)
    if formato == 'csv' and tipo == 'doacoes':
        return _resposta_csv(report, tipo, hoje)

    return render(request, 'sistema/relatorios.html', {
        'usuario': usuario,
        'form': form,
        'report': report,
        'report_types': RELATORIOS,
        'is_dashboard': tipo == 'indicadores',
    })
